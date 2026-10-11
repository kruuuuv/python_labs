"""CSV → XLSX через openpyxl."""

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter


def csv_to_xlsx(csv_path: str, xlsx_path: str) -> None:
    """CSV → XLSX.

    Первая строка CSV — заголовок. Лист "Sheet1".
    Автоширина колонок (минимум 8 символов).
    """
    p = Path(csv_path)
    if not p.exists():
        raise FileNotFoundError(f"Файл не найден: {csv_path}")

    with p.open(encoding="utf-8", newline="") as f:
        rows = list(csv.reader(f))

    if not rows:
        raise ValueError("Пустой CSV")

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    for row in rows:
        ws.append(row)

    # автоширина колонок
    for col_idx, col in enumerate(zip(*rows), start=1):
        max_len = max((len(str(cell)) for cell in col), default=0)
        ws.column_dimensions[get_column_letter(col_idx)].width = max(8, max_len + 2)

    out = Path(xlsx_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
